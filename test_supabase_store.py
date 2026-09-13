import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import openpyxl
from openpyxl.workbook.defined_name import DefinedName

SCRIPTS = Path(__file__).parent/'plugins/tohon-keiyakusho/skills/tohon-nyuryoku/scripts'
sys.path.insert(0,str(SCRIPTS))
import supabase_store as store
import store_and_fill
import fill_tohon


def sample():
    return {'一棟の建物':{'所在_市区町村町名':'試験市試験町','所在_番':'1','所在_枝番':'2','構造':'鉄筋コンクリート造','各階床面積':[{'階':'1階','面積':100}]},
        '専有部分':{'家屋番号_枝番':'3','建物の名称':'301','種類':'居宅','構造':'鉄筋コンクリート造1階建','階部分':'3','床面積':50,'建築時期':{'元号':'平成','年':10,'月':1,'日':1}},
        '所有者':{'氏名':'検証用法人','住所':'検証用住所'},'土地':[{'所在':'','地番':'','地目':'宅地','地積':100,'権利の種類':'所有権','持分_分子':1,'持分_分母':10}], '敷地権':True,'_要確認':[]}


class StoreTests(unittest.TestCase):
    def test_source_immutable(self):
        data=sample(); before=copy.deepcopy(data)
        store.convert(data,house_number='試験町1番2の3')
        self.assertEqual(data,before)

    def test_actual_keys(self):
        result=store.convert(sample(),house_number='試験町1番2の3')
        self.assertEqual(result['unit']['unit_floor'],'3階部分')
        self.assertEqual(result['unit']['built_date'],'1998-01-01')
        self.assertEqual(result['building']['land_lots'][0]['lot_number'],'1番2')

    def test_missing_mortgages_not_empty(self):
        self.assertNotIn('active_mortgages',store.convert(sample(),house_number='試験町1番2の3')['unit'])

    def test_explicit_mortgages(self):
        self.assertEqual(store.convert(sample(),house_number='試験町1番2の3',active_mortgages=[])['unit']['active_mortgages'],[])

    def test_no_invented_metadata(self):
        self.assertTrue(all('confidence' not in x for x in store.convert(sample(),house_number='試験町1番2の3')['values']))

    def test_metadata_passed(self):
        result=store.convert(sample(),house_number='試験町1番2の3',evidence={'所有者.氏名':{'confidence':0.9,'source_text':'確認済み原文','page_no':2}})
        self.assertEqual(next(x for x in result['values'] if x['field_code']=='current_owner_name')['page_no'],2)

    def test_house_required(self):
        with self.assertRaises(store.StoreError):store.convert(sample(),house_number='')
        with self.assertRaises(store.StoreError):store.convert(sample(),house_number='試験町1番2の4')

    def test_multi_share_not_collapsed(self):
        data=sample(); data['土地'].append(dict(data['土地'][0],持分_分子=2))
        result=store.convert(data,house_number='試験町1番2の3')
        self.assertNotIn('land_right_numerator',result['unit'])
        self.assertEqual(len(result['building']['land_lots']),2)

    def test_missing_owner_not_null_update(self):
        data=sample(); data['所有者']['氏名']=None
        self.assertNotIn('current_owner_name',store.convert(data,house_number='試験町1番2の3')['unit'])

    def test_idempotency_hash(self):
        kw={'house_number':'試験町1番2の3','file_hash':'a'*64,'original_filename':'test.pdf'}
        a=store.make_payload(sample(),**kw); b=store.make_payload(sample(),**kw)
        self.assertEqual(a['content_hash'],b['content_hash'])
        data=sample();data['所有者']['氏名']='別の検証用法人'
        self.assertNotEqual(a['content_hash'],store.make_payload(data,**kw)['content_hash'])
        self.assertEqual(a['document']['source_type'],'seller_provided')
        self.assertEqual(a['document']['status'],'provisional')

    def test_missing_credentials(self):
        with patch.object(store,'env',return_value=None),self.assertRaises(store.StoreError):store.rpc({})

    def test_excel_failure_gate_and_equivalence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); template=root/'template.xlsm'; src=root/'data.json'; output=root/'output.xlsm'
            src.write_text(json.dumps(sample(),ensure_ascii=False),encoding='utf8')
            wb=openpyxl.Workbook(); ws=wb.active;ws.title='基本入力';ws['A1']='=1+1';ws.merge_cells('B2:C2');ws['B2']='維持'
            wb.defined_names.add(DefinedName('謄本_所有者_氏名',attr_text="'基本入力'!$D$2"));wb.save(template);wb.close()
            with patch.object(store_and_fill,'save_to_supabase',side_effect=store.StoreError('失敗')),patch.object(store_and_fill.subprocess,'run') as runner:
                with self.assertRaises(store.StoreError):store_and_fill.run(src,template,output,house_number='試験町1番2の3',source_path=src)
                runner.assert_not_called();self.assertFalse(output.exists())
            baseline=root/'baseline.xlsm'
            import subprocess
            subprocess.run([sys.executable,str(SCRIPTS/'fill_tohon.py'),str(src),str(template),'-o',str(baseline)],stdout=subprocess.DEVNULL,check=False)
            with patch.object(store_and_fill,'save_to_supabase',return_value={'diffs_created':0}):
                store_and_fill.run(src,template,output,house_number='試験町1番2の3',source_path=src)
            with ZipFile(baseline) as a,ZipFile(output) as b:
                for name in a.namelist():
                    if name!='docProps/core.xml':self.assertEqual(a.read(name),b.read(name),name)


if __name__=='__main__':unittest.main()
