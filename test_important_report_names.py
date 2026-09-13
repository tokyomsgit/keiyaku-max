import copy
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook,load_workbook
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0,str(Path(__file__).parent/'plugins/tohon-keiyakusho/skills/tohon-nyuryoku/scripts'))
from add_important_report_names import add_names
from important_report_store import write_named_excel


class NamesTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.src=self.root/'source.xlsx';self.named=self.root/'named.xlsx'
        w=Workbook();s=w.active;s.title='基本入力';s['X13']='物件名';s['AV13']='号室'
        s['AB13']='旧名称';s['AS13']='●';s.merge_cells('AB13:AR14');s.merge_cells('AS13:AU14')
        s['AB13'].number_format='@';s['J26']='=AB13';w.create_sheet('他シート')['A1']='keep'
        w.defined_names.add(DefinedName('謄本_試験',attr_text="'基本入力'!$J$26"))
        w.save(self.src);w.close();add_names(self.src,self.named)
        self.fields={k:{'value':v,'approved':True,'confidence':.99,'needs_review':False,'page_no':1,'source_text':str(v)}
            for k,v in [('building_name','試験マンション'),('unit_name','0402号室')]}

    def tearDown(self):self.tmp.cleanup()

    def test_only_names_added(self):
        with zipfile.ZipFile(self.src) as a,zipfile.ZipFile(self.named) as b:
            self.assertEqual([n for n in a.namelist() if a.read(n)!=b.read(n)],['xl/workbook.xml'])
        w=load_workbook(self.named);self.assertEqual(w.defined_names['謄本_試験'].attr_text,"'基本入力'!$J$26");w.close()

    def test_add_is_idempotent(self):
        out=self.root/'again.xlsx';r=add_names(self.named,out);self.assertEqual(r['added'],0)
        with zipfile.ZipFile(out) as a,zipfile.ZipFile(self.named) as b:
            self.assertTrue(all(a.read(n)==b.read(n) for n in a.namelist()))

    def test_room_suffix_and_readback(self):
        out=self.root/'filled.xlsx';r=write_named_excel(self.named,out,self.fields)
        self.assertEqual(len(r['written']),2)
        w=load_workbook(out);self.assertEqual(w['基本入力']['AS13'].value,'0402');self.assertEqual(w['基本入力']['J26'].value,'=AB13');w.close()

    def test_uncertain_even_if_approved_is_blocked(self):
        for changes in ({'needs_review':True},{'confidence':.8},{'confidence':None},{'source_text':None},{'page_no':None},{'approved':False}):
            with self.subTest(changes=changes):
                fields=copy.deepcopy(self.fields)
                for f in fields.values():f.update(changes)
                r=write_named_excel(self.named,self.root/'blocked.xlsx',fields)
                self.assertEqual(r['written'],[]);self.assertFalse((self.root/'blocked.xlsx').exists())

    def test_blank_input_preserves_adjacent_cells(self):
        w=load_workbook(self.src);w['基本入力']['AB13']=None;w.save(self.src);w.close()
        blank=self.root/'blank.xlsx';add_names(self.src,blank)
        out=self.root/'filled.xlsx';r=write_named_excel(blank,out,self.fields)
        self.assertEqual(len(r['written']),2)
        w=load_workbook(out)
        self.assertEqual(w['基本入力']['AS13'].value,'0402')
        self.assertEqual(w['基本入力']['AV13'].value,'号室')
        self.assertEqual(w['基本入力']['J26'].value,'=AB13');w.close()

    def test_export_uses_names_after_cells_move(self):
        moved=self.root/'moved.xlsx'
        with zipfile.ZipFile(self.named) as a,zipfile.ZipFile(moved,'w') as b:
            for n in a.infolist():
                data=a.read(n.filename)
                if n.filename=='xl/workbook.xml':data=data.replace(b'$AB$13',b'$AS$13')
                b.writestr(n,data)
        r=write_named_excel(moved,self.root/'filled.xlsx',{'building_name':self.fields['building_name']})
        self.assertEqual(r['written'][0]['cell'],'AS13')


if __name__=='__main__':unittest.main()
