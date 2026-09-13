"""重調保存・明示承認・名前定義によるExcel反映。謄本処理は呼び変えない。"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import re
import urllib.request
import uuid
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape

from supabase_store import env, canonical, NoRedirect, StoreError
from important_report_schema import MAPPING, normalize


def rpc(name, payload):
    if name not in ('save_important_report','approve_important_report','get_approved_important_report'):
        raise StoreError('操作を確認してください。')
    url,key = env('SUPABASE_URL'),env('SUPABASE_SERVICE_ROLE_KEY')
    if not url or not key: raise StoreError('初期設定が完了していません。管理者へ連絡してください。')
    if not re.fullmatch(r'https://[a-z0-9]+\.supabase\.co/?',url): raise StoreError('接続先を確認してください。')
    request = urllib.request.Request(url.rstrip('/')+'/rest/v1/rpc/'+name,
        data=canonical({'p':payload}).encode('utf8'),headers={'apikey':key,'Authorization':'Bearer '+key,'Content-Type':'application/json'})
    try:
        with urllib.request.build_opener(NoRedirect()).open(request,timeout=60) as response: result=json.load(response)
        if not isinstance(result,dict) or not result.get('document_version_id'): raise ValueError()
        return result
    except (OSError,ValueError):
        raise StoreError('保存・承認に失敗しました。物件の対応、根拠、基準日を確認してください。Excel作成を停止しました。') from None


def make_payload(data, unit_id=None, house_number=None, source_type='seller_provided', status='provisional'):
    # Revalidate typed values without dropping evidence-validation warnings from the reader.
    checked=normalize(data)
    for k,item in checked['fields'].items():
        previous=data.get('fields',{}).get(k,{})
        item['review_reasons']=sorted(set(item['review_reasons']+previous.get('review_reasons',[])))
        item['needs_review']=item['needs_review'] or previous.get('needs_review',False)
    if source_type not in ('seller_provided','company_obtained') or status not in ('provisional','confirmed'):
        raise StoreError('資料の取得元・状態を確認してください。')
    if unit_id: unit_id=str(uuid.UUID(unit_id))
    source={k:data.get('source',{}).get(k) for k in ('file_hash','original_filename','storage_path')}
    if not re.fullmatch('[0-9a-f]{64}',source['file_hash'] or ''): raise StoreError('原本ハッシュがありません。')
    def val(k):
        item=checked['fields'][k]
        return None if item['needs_review'] else item['value']
    ident={'unit_id':unit_id,'house_number':house_number or checked.get('house_number'),
        'building_name':val('building_name'),'unit_name':val('unit_name'),'display_address':val('display_address')}
    if ident['unit_name']:
        ident['unit_name']=re.sub(r'\s*号室$', '', str(ident['unit_name'])).strip()
    checked.pop('source',None)
    payload={'identity':ident,'source':source,'source_type':source_type,'status':status,'raw_json':checked}
    # Moving the same original or replaying cache must not create a new version.
    identity_material=copy.deepcopy(payload);identity_material['source']={'file_hash':source['file_hash']}
    payload['import_key']=hashlib.sha256(canonical(identity_material).encode('utf8')).hexdigest()
    return payload


def save_to_supabase(data, **kwargs):
    return rpc('save_important_report',make_payload(data,**kwargs))


def approve(version_id, field_codes, reviewed_by, verified_against_original=False):
    if not field_codes or not set(field_codes)<=MAPPING.keys(): raise StoreError('承認項目を確認してください。')
    return rpc('approve_important_report',{'document_version_id':str(uuid.UUID(version_id)),
        'field_codes':field_codes,'reviewed_by':reviewed_by,'verified_against_original':bool(verified_against_original)})


def write_named_excel(template, output, approved_fields, mapping=None):
    """ZIP中の入力セル内容だけを置換し、その他のXMLとVBAを保持する。"""
    from openpyxl import load_workbook
    template,output=Path(template).resolve(),Path(output).resolve()
    if template==output or output.exists(): raise StoreError('別の出力先を指定してください。既存ファイルは上書きしません。')
    mapping=mapping or MAPPING
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    written=[]; skipped=[]; patches={}
    with zipfile.ZipFile(template) as z:
        wbxml=ET.fromstring(z.read('xl/workbook.xml'))
        rels={r.attrib['Id']:r.attrib['Target'] for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        sheets={s.attrib['name']:rels[s.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']] for s in wbxml.find('m:sheets',ns)}
        definitions={}
        for n in wbxml.findall('m:definedNames/m:definedName',ns):
            definitions.setdefault(n.attrib['name'],[]).append(n.text)
        wb=load_workbook(template,keep_vba=True,data_only=False)
        try:
            for code,item in approved_fields.items():
                name=mapping.get(code,{}).get('excel_named_range')
                if not item.get('approved') or item.get('value') is None or not name or len(definitions.get(name,[]))!=1:
                    skipped.append(code);continue
                ref=definitions[name][0]
                match=re.fullmatch(r"(?:'((?:[^']|'')+)'|([^'!]+))!\$?([A-Z]+)\$?(\d+)",ref or '')
                if not match: skipped.append(code);continue
                sheet=(match[1] or match[2]).replace("''", "'"); cell=match[3]+match[4]
                if sheet!='基本入力' or wb[sheet][cell].data_type=='f': skipped.append(code);continue
                if any(cell in r and cell!=r.start_cell.coordinate for r in wb[sheet].merged_cells.ranges): skipped.append(code);continue
                target=sheets[sheet]; target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
                data=patches.get(target,z.read(target).decode('utf8'))
                pattern=re.compile(r'<c\b(?=[^>]*\br="'+re.escape(cell)+r'")[^>]*(?:/>|>.*?</c>)',re.S)
                matches=list(pattern.finditer(data))
                if len(matches)!=1: skipped.append(code);continue
                node=matches[0].group(); opening=node[:node.index('>')+1].rstrip('/>')
                opening=re.sub(r'\s+t="[^"]*"','',opening)
                value=item['value']
                if isinstance(value,(int,float)) and not isinstance(value,bool): new=opening+'><v>'+str(value)+'</v></c>'
                else:
                    value=value if isinstance(value,str) else canonical(value)
                    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value): skipped.append(code);continue
                    new=opening+' t="inlineStr"><is><t xml:space="preserve">'+escape(value)+'</t></is></c>'
                data=data[:matches[0].start()]+new+data[matches[0].end():]
                patches[target]=data
                written.append({'field_code':code,'name':name,'sheet':sheet,'cell':cell,'value':value})
        finally: wb.close()
        if not written: return {'output':None,'written':[],'needs_review':skipped,'message':'対応する名前定義がないため、Excelへの反映はありません。'}
        output.parent.mkdir(parents=True,exist_ok=True)
        with zipfile.ZipFile(output,'x') as dst:
            for info in z.infolist(): dst.writestr(info,patches[info.filename].encode('utf8') if info.filename in patches else z.read(info.filename))
    check=load_workbook(output,keep_vba=True,data_only=False)
    try:
        if any(check[x['sheet']][x['cell']].value!=x['value'] for x in written): raise StoreError('Excelの再読込照合に失敗しました。')
    finally: check.close()
    return {'output':str(output),'written':written,'needs_review':skipped}


def export_approved(version_id,template,output):
    approved=rpc('get_approved_important_report',{'document_version_id':str(uuid.UUID(version_id))})
    result=write_named_excel(template,output,approved['fields'])
    result['unapproved_fields']=[k for k in MAPPING if k not in approved['fields']]
    return result


def main():
    parser=argparse.ArgumentParser(description='重要事項調査報告書の保存・承認・Excel反映')
    sub=parser.add_subparsers(dest='command',required=True)
    read=sub.add_parser('import');read.add_argument('pdf');read.add_argument('--cache-dir',default='output/important_report_cache')
    read.add_argument('--unit-id');read.add_argument('--house-number');read.add_argument('--source-type',default='seller_provided');read.add_argument('--status',default='provisional')
    adopt=sub.add_parser('approve');adopt.add_argument('version_id');adopt.add_argument('--fields',nargs='+',required=True);adopt.add_argument('--reviewed-by',required=True)
    adopt.add_argument('--verified-against-original',action='store_true',help='指定項目を原本で確認した場合のみ指定')
    export=sub.add_parser('export');export.add_argument('version_id');export.add_argument('template');export.add_argument('-o','--output',required=True)
    args=parser.parse_args()
    try:
        if args.command=='import':
            from important_report_reader import read_report
            data=read_report(args.pdf,args.cache_dir)
            result=save_to_supabase(data,unit_id=args.unit_id,house_number=args.house_number,source_type=args.source_type,status=args.status)
            result['要確認項目']=[MAPPING[k]['label'] for k,v in data['fields'].items() if v['needs_review'] or v['value'] is None]
        elif args.command=='approve': result=approve(args.version_id,args.fields,args.reviewed_by,args.verified_against_original)
        else: result=export_approved(args.version_id,args.template,args.output)
        print(json.dumps(result,ensure_ascii=False,indent=2))
    except (StoreError,ValueError,OSError):
        print('処理を完了できませんでした。資料・物件の対応・承認状態を管理者に確認してください。');return 1
    return 0


if __name__=='__main__': raise SystemExit(main())
