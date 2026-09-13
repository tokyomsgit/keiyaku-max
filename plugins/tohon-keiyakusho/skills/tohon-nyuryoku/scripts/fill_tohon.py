#!/usr/bin/env python3
"""
謄本の中間JSON を 契約書ひな型(.xlsm)の「基本入力」シートに書き込む。

使い方:
    python3 fill_tohon.py 謄本.json ひな型.xlsm -o 出力.xlsm

設計方針:
  - 書き込み先は mapping.json の対応表で決まる。このファイルには番地を書かない。
  - 対応表はひな型の「名前定義」を参照する。名前が無ければセル番地として解釈する。
  - 数式が入っているセルには絶対に書き込まない(スキップして警告)。
  - 結合セルは左上のアンカーにのみ書き込む。
  - マクロ(VBA)は保持する。
"""

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?([0-9]{1,7})$")
DEFINED_RE = re.compile(r"^'?([^'!]+)'?!\$?([A-Z]{1,3})\$?([0-9]{1,7})$")


class Report:
    """書き込み結果の記録。最後にまとめて人に見せる。"""

    def __init__(self):
        self.written = []
        self.skipped = []
        self.warnings = []

    def ok(self, cell, label, value):
        self.written.append((cell, label, value))

    def skip(self, cell, label, reason):
        self.skipped.append((cell, label, reason))

    def warn(self, message):
        if message not in self.warnings:
            self.warnings.append(message)

    def has_problems(self):
        return bool(self.skipped or self.warnings)

    def render(self):
        lines = [f"書き込み: {len(self.written)} 件"]
        for cell, label, value in self.written:
            shown = str(value)
            if len(shown) > 40:
                shown = shown[:40] + "…"
            lines.append(f"    {cell:<6} {label:<28} {shown}")

        if self.skipped:
            lines.append("")
            lines.append(f"スキップ: {len(self.skipped)} 件")
            for cell, label, reason in self.skipped:
                lines.append(f"    {cell:<6} {label:<28} {reason}")

        if self.warnings:
            lines.append("")
            lines.append("要確認:")
            for w in self.warnings:
                lines.append(f"    ! {w}")

        return "\n".join(lines)


class Target:
    """書き込み先の解決役。名前定義とセル番地の両方を受け付ける。"""

    def __init__(self, workbook, sheet_name, report):
        self.wb = workbook
        self.sheet = workbook[sheet_name]
        self.report = report
        self.merge_index = {}
        for merged in self.sheet.merged_cells.ranges:
            anchor = (merged.min_col, merged.min_row)
            for row in self.sheet[str(merged)]:
                for cell in row:
                    self.merge_index[(cell.column, cell.row)] = anchor

    def locate(self, ref, label):
        """参照(名前定義またはセル番地)を (列番号, 行番号) に解決する。"""
        defined = self.wb.defined_names.get(ref)
        if defined is not None:
            m = DEFINED_RE.match(defined.attr_text.strip())
            if m:
                sheet_name, col, row = m.groups()
                if sheet_name != self.sheet.title:
                    self.report.warn(
                        f"{label}: 名前定義「{ref}」が別シート({sheet_name})を指しています。"
                    )
                    return None
                return column_index_from_string(col), int(row)
            self.report.warn(f"{label}: 名前定義「{ref}」の形式を解釈できません。")
            return None

        m = CELL_RE.match(ref.strip())
        if m:
            col, row = m.groups()
            self.report.warn(
                f"{label}: 「{ref}」は名前定義ではなくセル番地として扱いました。"
                "ひな型に名前定義を追加すると、行の増減に強くなります。"
            )
            return column_index_from_string(col), int(row)

        self.report.warn(f"{label}: 参照「{ref}」を解決できませんでした。書き込みません。")
        return None

    def write(self, ref, value, label, row_offset=0):
        """1セル書き込む。数式セルは絶対に壊さない。"""
        if is_blank(value):
            return
        pos = self.locate(ref, label)
        if pos is None:
            return
        col, row = pos[0], pos[1] + row_offset
        self.write_at(col, row, value, label)

    def write_at(self, col, row, value, label):
        anchor = self.merge_index.get((col, row), (col, row))
        coord = f"{get_column_letter(anchor[0])}{anchor[1]}"
        current = self.sheet.cell(row=anchor[1], column=anchor[0]).value

        if isinstance(current, str) and current.startswith("="):
            self.report.skip(coord, label, f"数式のため書き込まず({current[:24]})")
            return

        self.sheet.cell(row=anchor[1], column=anchor[0], value=value)
        self.report.ok(coord, label, value)

    def clear_at(self, col, row):
        anchor = self.merge_index.get((col, row), (col, row))
        cell = self.sheet.cell(row=anchor[1], column=anchor[0])
        if not (isinstance(cell.value, str) and cell.value.startswith("=")):
            cell.value = None

    def read(self, ref, label):
        pos = self.locate(ref, label)
        if pos is None:
            return None
        return self.sheet.cell(row=pos[1], column=pos[0]).value


def dig(data, dotted_path):
    """'専有部分.建築時期.年' のようなパスで JSON を辿る。無ければ None。"""
    node = data
    for key in dotted_path.split("."):
        if not isinstance(node, dict) or key not in node:
            return None
        node = node[key]
    return node


def is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def fill_single_fields(target, mapping, data):
    for path, ref in mapping["single"].items():
        if path.startswith("_"):
            continue
        target.write(ref, dig(data, path), path)


def fill_land(target, mapping, data, report):
    spec = mapping["land"]
    parcels = data.get("土地") or []

    if len(parcels) > spec["max_rows"]:
        report.warn(
            f"土地が {len(parcels)} 筆あります。ひな型は {spec['max_rows']} 筆までなので、"
            f"{spec['max_rows'] + 1} 筆目以降は手入力してください。"
        )

    for i, parcel in enumerate(parcels[: spec["max_rows"]]):
        offset = i * spec["step"]
        for field, ref in spec["head"].items():
            if field.startswith("_"):
                continue
            target.write(ref, parcel.get(field), f"土地{i + 1}.{field}", row_offset=offset)


def fill_floors(target, mapping, data, report):
    """各階の床面積を計算表に並べ、延床面積セルに合計の数式を入れる。"""
    spec = mapping["floors"]
    floors = dig(data, "一棟の建物.各階床面積") or []
    if not floors:
        return

    capacity = spec["rows"]
    if len(floors) > capacity:
        report.warn(
            f"各階床面積が {len(floors)} 行あり、計算表({capacity}行)に収まりません。超過分は手入力してください。"
        )

    used = 0
    for i, floor in enumerate(floors[:capacity]):
        area = floor.get("面積")
        if is_blank(area):
            continue
        label = floor.get("階", i + 1)
        # 階ラベルも書き換える。ひな型の初期値は 1,2,3… の固定なので、
        # 地下階のある建物ではラベルと面積がズレてしまうため。
        target.write(spec["floor_head"], label, f"各階.{label}(階名)", row_offset=i)
        target.write(spec["area_head"], area, f"各階.{label}(面積)", row_offset=i)
        used = i + 1

    # 使わなかった行に古い階番号が残ると誤読のもとなので消す。
    head = target.locate(spec["floor_head"], "各階(初期値の消去)")
    if head:
        for i in range(used, capacity):
            target.clear_at(head[0], head[1] + i)

    area_head = target.locate(spec["area_head"], "延床面積(合計)")
    if area_head:
        col = get_column_letter(area_head[0])
        first, last = area_head[1], area_head[1] + capacity - 1
        target.write(spec["total"], f"=SUM({col}{first}:{col}{last})", "延床面積(合計)")


def fill_shikichiken(target, mapping, data, report):
    spec = mapping["shikichiken"]
    flag = data.get("敷地権")
    if flag is None:
        report.warn("敷地権の有無が未指定です。ひな型の初期値(有)のままになっています。")
        return

    target.write(spec["yes"], spec["on"] if flag else spec["off"], "敷地権.有")
    target.write(spec["no"], spec["off"] if flag else spec["on"], "敷地権.無")


def fill_shakuchi(target, mapping, data, report):
    """借地権付き物件のみ。JSONに「借地」が無ければ何もしない。"""
    spec = mapping.get("shakuchi")
    shakuchi = data.get("借地")
    if not spec or not shakuchi:
        return

    for path, ref in spec["fields"].items():
        if path.startswith("_"):
            continue
        target.write(ref, dig(data, path), path)

    checks = spec["checks"]
    on, off = spec["on"], spec["off"]

    for group in ["面積の根拠", "法区分", "新法の種類"]:
        chosen = shakuchi.get(group)
        for label, ref in checks[group].items():
            if label.startswith("_"):
                continue
            target.write(ref, on if chosen == label else off, f"借地.{group}.{label}")

    if shakuchi.get("譲渡承諾特約") is not None:
        target.write(
            checks["譲渡承諾特約"],
            on if shakuchi["譲渡承諾特約"] else off,
            "借地.譲渡承諾特約",
        )

    if is_blank(shakuchi.get("地代_月額")):
        report.warn(
            "借地の地代(月額)が空です。謄本が「3.3㎡当り月◯円」のような単価表記の場合、"
            "月額の総額は自動計算できません。地主または借地説明書で確認してください。"
        )


def check_jusho_hyoji(target, mapping, data, report):
    """
    ひな型は「所在」の市区町村・町名部分を住居表示から作る。
    ここが謄本と食い違うと、重説・売契の所在がまるごと間違う。
    """
    tohon_town = dig(data, "一棟の建物.所在_市区町村町名")
    spec = mapping.get("jusho_hyoji")
    if is_blank(tohon_town) or not spec:
        return

    ku = target.read(spec["ku"], "住居表示.区") or ""
    machi = target.read(spec["machi"], "住居表示.町名") or ""
    current = f"{ku}{machi}".strip().replace("\u3000", "")
    expected = str(tohon_town).strip().replace("\u3000", "")

    if current != expected:
        report.warn(
            f"住居表示は「{current}」ですが、謄本の所在は「{expected}」です。"
            "所在は住居表示から組み立てられるため、先に住居表示を入力してください。"
        )


def sanity_check(data, report):
    """人が見落としやすいところを機械で拾う。"""
    owner = dig(data, "所有者.氏名")
    if isinstance(owner, str) and any(sep in owner for sep in ["、", ",", "及び", "・"]):
        report.warn(
            f"所有者が複数名の可能性があります({owner})。ひな型は1名分の枠なので確認してください。"
        )

    senyu = dig(data, "専有部分.床面積")
    if senyu is not None and not isinstance(senyu, (int, float)):
        report.warn(f"専有部分の床面積が数値ではありません({senyu!r})。")

    for i, parcel in enumerate(data.get("土地") or []):
        bunbo, bunshi = parcel.get("持分_分母"), parcel.get("持分_分子")
        if isinstance(bunbo, (int, float)) and isinstance(bunshi, (int, float)):
            if bunshi > bunbo:
                report.warn(f"土地{i + 1}: 持分の分子が分母より大きいです({bunshi}/{bunbo})。")

    if not dig(data, "一棟の建物.所在_番"):
        report.warn("一棟の建物の所在(番)が空です。謄本から拾えていない可能性があります。")


def count_formulas(path):
    wb = openpyxl.load_workbook(path, keep_vba=True)
    total = sum(
        1
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    wb.close()
    return total


def main():
    parser = argparse.ArgumentParser(description="謄本の中間JSONを契約書ひな型に書き込みます。")
    parser.add_argument("json_path", help="謄本の中間JSON")
    parser.add_argument("template_path", help="契約書ひな型(.xlsm)")
    parser.add_argument("-o", "--output", required=True, help="出力先(.xlsm)")
    parser.add_argument(
        "-m", "--mapping",
        default=str(Path(__file__).parent / "mapping.json"),
        help="対応表(既定: 同じフォルダの mapping.json)",
    )
    args = parser.parse_args()

    data = json.loads(Path(args.json_path).read_text(encoding="utf-8"))
    mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8"))

    before = count_formulas(args.template_path)

    shutil.copy(args.template_path, args.output)
    workbook = openpyxl.load_workbook(args.output, keep_vba=True)

    # 名前定義の無い古いひな型を誤って使わないための門番。
    sentinels = [mapping["single"]["所有者.氏名"], mapping["floors"]["total"]]
    if not any(s in workbook.defined_names for s in sentinels):
        workbook.close()
        Path(args.output).unlink(missing_ok=True)
        print(
            "このひな型には名前定義がありません。\n"
            "名前定義版のひな型(基本入力に「謄本_」で始まる名前が定義されているもの)を使ってください。\n"
            "古いひな型を使い続ける場合は、mapping.json の名前をセル番地に書き換える必要があります。"
        )
        return 1

    report = Report()
    target = Target(workbook, mapping["sheet"], report)

    sanity_check(data, report)
    check_jusho_hyoji(target, mapping, data, report)
    fill_single_fields(target, mapping, data)
    fill_land(target, mapping, data, report)
    fill_floors(target, mapping, data, report)
    fill_shikichiken(target, mapping, data, report)
    fill_shakuchi(target, mapping, data, report)

    workbook.save(args.output)
    workbook.close()

    after = count_formulas(args.output)
    if after < before:
        report.warn(
            f"数式が {before - after} 件減りました。ひな型が壊れている可能性があるので出力を使わないでください。"
        )

    print(report.render())
    print()
    print(f"出力: {args.output}")
    if report.has_problems():
        print("→ 上の「スキップ」「要確認」を見てから使ってください。")
        return 1
    print("→ Excelで開いて、謄本と並べて目視確認してください。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
