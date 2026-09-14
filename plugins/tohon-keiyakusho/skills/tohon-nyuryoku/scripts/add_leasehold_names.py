"""Add missing leasehold defined names to an official contract template copy."""
import argparse,json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName


def add_names(source,output):
    source,output=Path(source),Path(output)
    if source.resolve()==output.resolve():raise ValueError('原本は上書きできません。')
    mapping=json.loads(Path(__file__).with_name('leasehold_template_names.json').read_text(encoding='utf8'))
    wb=load_workbook(source,keep_vba=True,data_only=False)
    try:
        if '基本入力' not in wb.sheetnames or '借地' not in wb.sheetnames:raise ValueError('正式な契約書ひな形ではありません。')
        added=[]
        for name,target in mapping.items():
            current=wb.defined_names.get(name)
            if current and current.attr_text!=target:raise ValueError(f'既存の名前定義が異なります: {name}')
            if not current:wb.defined_names.add(DefinedName(name,attr_text=target));added.append(name)
        wb.save(output)
    finally:wb.close()
    check=load_workbook(output,read_only=True,data_only=False,keep_vba=True)
    try:
        if any(check.defined_names.get(k) is None for k in mapping):raise ValueError('名前定義の再読込に失敗しました。')
    finally:check.close()
    return added


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source');p.add_argument('-o','--output',required=True);a=p.parse_args()
    print(f'借地用の名前定義を{len(add_names(a.source,a.output))}件追加しました。')
