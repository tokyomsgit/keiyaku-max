"""正式ひな形への名前追加専用。アプリの転記は名前だけを参照する。"""
import argparse
import json
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape, quoteattr


def add_names(source, output, definitions=None):
    from openpyxl import load_workbook
    source,output=Path(source).resolve(),Path(output).resolve()
    if source==output or output.exists(): raise ValueError('入力と異なる未使用の出力先を指定してください。')
    definitions=definitions or json.loads(Path(__file__).with_name('important_report_template_names.json').read_text(encoding='utf8'))
    wb=load_workbook(source,read_only=False,data_only=False)
    try:
        for name,spec in definitions.items():
            m=re.fullmatch(r"'基本入力'!\$([A-Z]+)\$(\d+)",spec['reference'])
            if not m: raise ValueError('名前の参照先が未対応です。')
            ws=wb['基本入力'];cell=m[1]+m[2]
            if ws[spec['label_cell']].value!=spec['label'] or ws[cell].data_type=='f':
                raise ValueError('入力欄の配置が一致しません。')
            if any(cell in merged and cell!=merged.start_cell.coordinate for merged in ws.merged_cells.ranges):
                raise ValueError('結合セルの先頭ではありません。')
    finally: wb.close()
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(source) as src:
        raw=src.read('xl/workbook.xml')
        root=ET.fromstring(raw)
        names=root.findall('m:definedNames/m:definedName',ns)
        additions=[]
        for name,spec in definitions.items():
            found=[n for n in names if n.get('name')==name]
            if found:
                if len(found)!=1 or found[0].text!=spec['reference'] or found[0].get('localSheetId') is not None:
                    raise ValueError('既存の同名定義と参照先が異なります。')
                continue
            additions.append('<definedName name='+quoteattr(name)+'>'+escape(spec['reference'])+'</definedName>')
        token=b'</definedNames>'
        if raw.count(token)!=1: raise ValueError('名前定義の構造が未対応です。')
        changed=raw.replace(token,''.join(additions).encode('utf8')+token)
        output.parent.mkdir(parents=True,exist_ok=True)
        with zipfile.ZipFile(output,'x') as dst:
            for info in src.infolist():dst.writestr(info,changed if info.filename=='xl/workbook.xml' else src.read(info.filename))
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(output) as b:
        assert a.namelist()==b.namelist()
        assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='xl/workbook.xml')
    return {'added':len(additions),'total':len(names)+len(additions),'definitions':definitions}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description='セルを変更せず正式ひな形へ重調用名前定義を追加')
    parser.add_argument('template');parser.add_argument('output');a=parser.parse_args()
    print(json.dumps(add_names(a.template,a.output),ensure_ascii=False,indent=2))
