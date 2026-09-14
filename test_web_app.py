import copy
import http.client
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from web_app import make_server
from web_data import Workspace,StoreError,field_view


class WebTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        registry={'一棟の建物':{'所在_市区町村町名':'試験市一丁目','構造':'鉄筋コンクリート造'},
          '専有部分':{'家屋番号_枝番':101,'建物の名称':'101','床面積':40},'所有者':{'氏名':'試験所有者','住所':'試験市'},'土地':[]}
        fields={k:{'value':v,'source_text':str(v),'page_no':1,'confidence':.99,'needs_review':False}
          for k,v in [('building_name','試験マンション'),('unit_name','101号室'),('management_fee',12000)]}
        (self.root/'registry.json').write_text(json.dumps(registry),encoding='utf8')
        (self.root/'report.json').write_text(json.dumps({'fields':fields}),encoding='utf8')
        self.env=patch.dict(os.environ,{'DEMO_MODE':'true','DEMO_WRITE_DB':'false','DEMO_REGISTRY_JSON':str(self.root/'registry.json'),
          'DEMO_REPORT_JSON':str(self.root/'report.json'),'DEMO_EVIDENCE_JSON':'','DEMO_HOUSE_NUMBER':'試験一丁目1の101',
          'DEMO_APPROVED_FIELDS':'building_name,unit_name','WEB_OUTPUT_DIR':str(self.root/'output')})
        self.env.start();self.load=patch('web_data.load_env');self.load.start()
        self.network=patch('urllib.request.build_opener',side_effect=AssertionError('Network forbidden in demo'));self.network.start()
        self.w=Workspace()

    def tearDown(self):
        self.network.stop();self.load.stop();self.env.stop();self.tmp.cleanup()

    def test_demo_has_documents_and_zero_network(self):
        s=self.w.public();self.assertEqual(s['usage'],{'ai_calls':0,'db_reads':0,'db_writes':0})
        self.assertEqual([d['type'] for d in s['cases'][0]['documents']],['registry','important_report'])
        self.assertNotIn('registry',s);self.assertNotIn('SUPABASE_SERVICE_ROLE_KEY',json.dumps(s))

    def test_hold_does_not_change_master(self):
        before=copy.deepcopy(self.w.raw['demo-case']['unit'])
        self.w.decide('demo-case','demo-current_owner_name','hold')
        self.assertEqual(before,self.w.raw['demo-case']['unit'])
        self.assertEqual(self.w.cases[0]['diffs'][1]['review_status'],'reviewed')

    def test_explicit_adopt_and_duplicate(self):
        self.w.decide('demo-case','demo-management_fee','adopt')
        self.assertEqual(self.w.raw['demo-case']['unit']['management_fee'],13500)
        with self.assertRaises(StoreError):self.w.decide('demo-case','demo-management_fee','adopt')

    def test_stale_master_and_untrusted_adoption_blocked(self):
        self.w.raw['demo-case']['unit']['management_fee']=1
        with self.assertRaises(StoreError):self.w.decide('demo-case','demo-management_fee','adopt')
        self.w.cases[0]['diffs'][1]['can_adopt']=False
        with self.assertRaises(StoreError):self.w.decide('demo-case','demo-current_owner_name','adopt')

    def test_missing_evidence_and_confidence_are_review(self):
        for changes in ({'confidence':None},{'confidence':.5},{'page_no':None},{'source_text':None},{'needs_review':True}):
            f={'value':1,'confidence':.99,'source_text':'1','page_no':1,**changes}
            self.assertTrue(field_view('management_fee',f)['needs_review'])

    def test_http_host_csrf_and_private_paths(self):
        server=make_server(self.w,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        port=server.server_port
        def request(method,path,body=None,headers=None):
            c=http.client.HTTPConnection('127.0.0.1',port,timeout=5)
            c.request(method,path,body,headers or {});r=c.getresponse();result=(r.status,r.read(),dict(r.getheaders()));c.close();return result
        try:
            status,body,_=request('GET','/api/state');self.assertEqual(status,200);csrf=json.loads(body)['csrf']
            for sheet in ('theme','base','layout','components'):
                self.assertEqual(request('GET','/styles/'+sheet+'.css')[0],200)
            for path in ('/.env','/web_data.py','/../.env','/download/unknown'):
                self.assertEqual(request('GET',path)[0],404)
            self.assertEqual(request('GET','/api/state',headers={'Host':'evil.example'})[0],403)
            data=json.dumps({'case_id':'demo-case','diff_id':'demo-management_fee','action':'adopt'})
            self.assertEqual(request('POST','/api/decision',data)[0],403)
            headers={'Origin':f'http://127.0.0.1:{port}','X-CSRF-Token':csrf,'Content-Type':'application/json'}
            self.assertEqual(request('POST','/api/decision',data,headers)[0],200)
            self.assertEqual(request('POST','/api/decision',data,{**headers,'Origin':'https://evil.example'})[0],403)
            path=self.root/'download.xlsm';path.write_bytes(b'only-generated-file');self.w.files['known']=path
            code,body,h=request('GET','/download/known');self.assertEqual(body,path.read_bytes());self.assertEqual(code,200)
            self.assertIn('attachment',h['Content-Disposition'])
        finally:server.shutdown();server.server_close();thread.join()

    def test_generation_uses_adopted_master_and_explicit_report_baseline(self):
        from openpyxl import Workbook
        from openpyxl.workbook.defined_name import DefinedName
        import subprocess
        def fill(args,**kwargs):
            incoming=json.loads(Path(args[4]).read_text(encoding='utf8'))
            self.assertEqual(incoming['所有者']['氏名'],'デモ用変更候補（実在しない法人）')
            w=Workbook();w.active.title='基本入力'
            w.defined_names.add(DefinedName('重調_building_name',attr_text="'基本入力'!$A$1"))
            w.defined_names.add(DefinedName('重調_unit_name',attr_text="'基本入力'!$B$1"))
            w.active['A1']='旧名称';w.active['B1']='旧号室';w.active['C1']='=1+1';w.save(args[-1]);w.close()
            return subprocess.CompletedProcess(args,0,'','')
        self.w.decide('demo-case','demo-current_owner_name','adopt')
        def preserve(template,candidate,output):
            import shutil
            shutil.copyfile(candidate,output);return []
        with patch('web_data.subprocess.run',side_effect=fill),patch('web_excel.preserve_template',side_effect=preserve):result=self.w.generate('demo-case')
        self.assertEqual(result['report_written'],2)
        self.assertTrue(next(iter(self.w.files.values())).exists())
        self.assertEqual(self.w.public()['usage']['db_writes'],0)

    def test_legacy_changes_to_formulas_styles_and_other_sheets_are_discarded(self):
        from openpyxl import Workbook,load_workbook
        from openpyxl.styles import Font
        from web_excel import preserve_template
        import zipfile
        w=Workbook();s=w.active;s.title='基本入力';s['A1']='old';s['A2']='=1+1';s['A3']=None;s['A3'].font=Font(bold=True)
        w.create_sheet('他シート')['A1']='preserve';src=self.root/'source.xlsx';candidate=self.root/'candidate.xlsx';out=self.root/'out.xlsx';w.save(src)
        s['A1']='new & value';s['A1'].font=Font(italic=True);s['A2']='=9+9';s['A3']=123;w['他シート']['A1']=None;w.save(candidate);w.close()
        rows=preserve_template(src,candidate,out);self.assertEqual(len(rows),2)
        w=load_workbook(out);self.assertEqual(w['基本入力']['A1'].value,'new & value');self.assertFalse(w['基本入力']['A1'].font.italic)
        self.assertEqual(w['基本入力']['A2'].value,'=1+1');self.assertEqual(w['基本入力']['A3'].value,123);self.assertEqual(w['他シート']['A1'].value,'preserve');w.close()
        with zipfile.ZipFile(src) as a,zipfile.ZipFile(out) as b:
            self.assertEqual(set(n for n in a.namelist() if a.read(n)!=b.read(n)),{'xl/worksheets/sheet1.xml','xl/workbook.xml'})


if __name__=='__main__':unittest.main()
