import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

SCRIPTS=Path(__file__).parent/'plugins/tohon-keiyakusho/skills/tohon-nyuryoku/scripts'
sys.path.insert(0,str(SCRIPTS))
from important_report_schema import normalize, FIELDS
from important_report_store import make_payload,write_named_excel,export_approved,StoreError


def report():
    return {'is_important_report':True,'house_number':None,'issuer_name':'試験管理株式会社',
        'fields':{'management_fee':{'value':'１２，０００円','source_label':'月額管理費','source_section':'住戸',
        'source_text':'月額管理費１２，０００円','page_no':2,'value_as_of_date':'2026-01-31','confidence':.99}},
        'source':{'file_hash':'a'*64,'original_filename':'synthetic.pdf','storage_path':'private/synthetic.pdf'}}


class ImportantReportTests(unittest.TestCase):
    def test_standard_fields_and_number(self):
        d=normalize(report());self.assertEqual(len(d['fields']),35)
        self.assertEqual(d['fields']['management_fee']['value'],12000)
        self.assertIsNone(d['fields']['repair_reserve_fee']['value'])

    def test_field_dates_not_document_date(self):
        r=report();r['as_of_date']='2026-09-01'
        r['fields']['repair_reserve_fee']=dict(r['fields']['management_fee'],value_as_of_date=None)
        d=normalize(r);self.assertEqual(d['fields']['management_fee']['value_as_of_date'],'2026-01-31')
        self.assertIsNone(d['fields']['repair_reserve_fee']['value_as_of_date'])

    def test_partial_date_is_not_first_day(self):
        r=report();r['fields']['management_fee']['value_as_of_date']='2026-01'
        self.assertIsNone(normalize(r)['fields']['management_fee']['value_as_of_date'])

    def test_page_correction_and_scan_review(self):
        r=report();pages=[{'page_no':1,'mode':'text','text':'月額管理費１２，０００円'}, {'page_no':2,'mode':'text','text':'他の文章'}]
        self.assertEqual(normalize(r,pages)['fields']['management_fee']['page_no'],1)
        pages[1]['mode']='scan'
        self.assertTrue(normalize(r,pages)['fields']['management_fee']['needs_review'])

    def test_conflicting_amount_blocked(self):
        r=report();r['fields']['management_fee']['value']='13000'
        self.assertTrue(normalize(r)['fields']['management_fee']['needs_review'])

    def test_non_report_rejected(self):
        r=report();r['is_important_report']=False
        with self.assertRaises(ValueError): normalize(r)

    def test_duplicate_material_excludes_cache_and_path(self):
        r=report();a=make_payload(r)
        r['source']['cache_reused']=True;r['source']['storage_path']='moved/synthetic.pdf'
        self.assertEqual(a['import_key'],make_payload(r)['import_key'])
        r['fields']['management_fee']['value']='13000'
        self.assertNotEqual(a['import_key'],make_payload(r)['import_key'])

    def test_review_reasons_preserved_on_save(self):
        r=normalize(report());r['fields']['management_fee'].update(needs_review=True,review_reasons=['画像の確認'])
        self.assertTrue(make_payload(r)['raw_json']['fields']['management_fee']['needs_review'])

    def test_failed_db_never_writes_excel(self):
        with patch('important_report_store.rpc',side_effect=StoreError('failed')), patch('important_report_store.write_named_excel') as w:
            with self.assertRaises(StoreError): export_approved('00000000-0000-0000-0000-000000000001','unused','unused2')
            w.assert_not_called()

    def test_named_write_formula_and_package_preservation(self):
        from openpyxl import Workbook,load_workbook
        from openpyxl.workbook.defined_name import DefinedName
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);src=root/'input.xlsx';out=root/'output.xlsx'
            w=Workbook();ws=w.active;ws.title='基本入力';ws['B2']=5;ws['C2']='=B2*2';ws.merge_cells('B2:B3')
            ws['B2'].number_format='#,##0';w.create_sheet('他シート')['A1']='unchanged'
            w.defined_names.add(DefinedName('試験管理費',attr_text="'基本入力'!$B$2"))
            w.defined_names.add(DefinedName('試験数式',attr_text="'基本入力'!$C$2"));w.save(src)
            fields={'management_fee':{'value':12000,'approved':True},'repair_reserve_fee':{'value':9,'approved':True},'unit_name':{'value':'=cmd','approved':False}}

            for item in fields.values():item.update(confidence=.99,needs_review=False,source_text='確認済み試験値',page_no=1)
            mapping={'management_fee':{'excel_named_range':'試験管理費'},'repair_reserve_fee':{'excel_named_range':'試験数式'}}
            r=write_named_excel(src,out,fields,mapping);self.assertEqual(len(r['written']),1)
            c=load_workbook(out);self.assertEqual(c['基本入力']['B2'].value,12000);self.assertEqual(c['基本入力']['C2'].value,'=B2*2');c.close()
            with zipfile.ZipFile(src) as a,zipfile.ZipFile(out) as b:
                self.assertEqual(set(x for x in a.namelist() if a.read(x)!=b.read(x)),{'xl/worksheets/sheet1.xml','xl/workbook.xml'})
            with self.assertRaises(StoreError):write_named_excel(src,out,fields,mapping)

    def test_unmapped_does_not_create_contract(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/'template.xlsx';out=Path(tmp)/'output.xlsx';Workbook().save(src)
            result=write_named_excel(src,out,{'management_fee':{'value':12000,'approved':True}})
            self.assertIsNone(result['output']);self.assertFalse(out.exists())


if __name__=='__main__':unittest.main()
