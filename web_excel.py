"""Apply legacy writer values to the original package, preserving its other parts."""
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
from openpyxl import load_workbook


def preserve_template(template,candidate,output):
    from supabase_store import StoreError
    a=load_workbook(template,data_only=False);b=load_workbook(candidate,data_only=False)
    try:
        changes=[]
        for row in b['基本入力']:
            for c in row:
                old=a['基本入力'][c.coordinate]
                if old.data_type=='f' or c.data_type=='f':continue
                if old.value!=c.value:changes.append((c.coordinate,c.value))
    finally:a.close();b.close()
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(template) as src:
        wb=ET.fromstring(src.read('xl/workbook.xml'))
        rels={x.attrib['Id']:x.attrib['Target'] for x in ET.fromstring(src.read('xl/_rels/workbook.xml.rels'))}
        sheet=next(s for s in wb.find('m:sheets',ns) if s.attrib['name']=='基本入力')
        target=rels[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
        target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        xml=src.read(target).decode('utf8')
        for coord,value in changes:
            pattern=re.compile(r'<c\b(?=[^>]*\br="'+re.escape(coord)+r'")[^>]*?(?:/>|>.*?</c>)',re.S)
            found=list(pattern.finditer(xml))
            if len(found)!=1:raise StoreError('ひな形の入力セルを確認してください。')
            original=found[0].group();attrs=re.match(r'<c\b([^>]*?)(?:/?>)',original)[1]
            attrs=re.sub(r'\s+t="[^"]*"','',attrs)
            if value is None:body=''
            elif isinstance(value,bool):attrs+=' t="b"';body='<v>'+str(int(value))+'</v>'
            elif isinstance(value,(int,float)):body='<v>'+str(value)+'</v>'
            elif isinstance(value,str):
                attrs+=' t="inlineStr"';body='<is><t xml:space="preserve">'+escape(value)+'</t></is>'
            else:raise StoreError('転記する値の形式を確認してください。')
            xml=pattern.sub(lambda _: '<c'+attrs+'>'+body+'</c>',xml,count=1)
        workbook=src.read('xl/workbook.xml').decode('utf8')
        calc=re.search(r'<calcPr\b[^>]*/>',workbook)
        if not calc:raise StoreError('ひな形の再計算設定を確認してください。')
        updated=calc.group()
        for attr in ('fullCalcOnLoad','forceFullCalc','calcMode'):updated=re.sub(r'\s+'+attr+r'="[^"]*"','',updated)
        updated=updated[:-2]+' fullCalcOnLoad="1" forceFullCalc="1" calcMode="auto"/>'
        workbook=workbook[:calc.start()]+updated+workbook[calc.end():]
        patches={target:xml.encode('utf8'),'xl/workbook.xml':workbook.encode('utf8')}
        with zipfile.ZipFile(output,'x',zipfile.ZIP_DEFLATED) as out:
            for item in src.infolist():out.writestr(item,patches.get(item.filename,src.read(item.filename)))
    check=load_workbook(output,read_only=True,data_only=False)
    try:
        if any(check['基本入力'][coord].value!=value for coord,value in changes):raise StoreError('転記値の再読込が一致しません。')
    finally:check.close()
    return [{'cell':coord,'value':value} for coord,value in changes]
