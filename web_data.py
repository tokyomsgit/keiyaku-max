"""Local web workspace. This module never imports or invokes an AI reader."""
import copy
import datetime as dt
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
import urllib.request
import uuid

ROOT=Path(__file__).resolve().parent
SCRIPTS=ROOT/'plugins/tohon-keiyakusho/skills/tohon-nyuryoku/scripts'
sys.path.insert(0,str(SCRIPTS))
from supabase_store import env,canonical,NoRedirect,StoreError,convert
from important_report_schema import MAPPING
from important_report_store import write_named_excel
from management_rules_schema import MAPPING as RULES_MAPPING
from purchase_excel import MAPPING as PURCHASE_MAPPING


def load_env():
    p=ROOT/'.env'
    if p.exists():
        for line in p.read_text(encoding='utf-8-sig').splitlines():
            if line.strip() and not line.lstrip().startswith('#') and '=' in line:
                k,v=line.split('=',1)
                os.environ.setdefault(k.strip(),v.strip().strip('"').strip("'"))


def flag(name,default=False):return str(env(name) or str(default)).lower()=='true'


def read_json(path):return json.loads(Path(path).read_text(encoding='utf-8-sig'))


LABELS={**{k:v['label'] for k,v in MAPPING.items()},'registry_location':'登記上の所在','building_structure':'一棟の構造',
 'floor_areas':'各階床面積','land_lots':'土地','house_number':'家屋番号','unit_type':'建物の種類','unit_structure':'専有部分の構造',
 'unit_floor':'所在階','registered_area':'登記面積','built_date':'新築年月日','current_owner_name':'現在の所有者',
 'current_owner_address':'所有者住所','has_land_right':'敷地権の有無','land_right_type':'敷地権の種類',
 'land_right_numerator':'敷地権割合（分子）','land_right_denominator':'敷地権割合（分母）','active_mortgages':'有効な抵当権'}


def evidence_item(raw,code):
    paths={'building_name':'building.name','registry_location':'building.location','building_structure':'building.structure',
      'house_number':'unit.house_number','unit_name':'unit.name','unit_type':'unit.type','unit_structure':'unit.structure',
      'unit_floor':'unit.floor','registered_area':'unit.registered_area','built_date':'unit.built_date',
      'current_owner_name':'owner.name','current_owner_address':'owner.address','has_land_right':'land_right.exists'}
    v=raw
    for p in paths.get(code,'').split('.'):v=v.get(p,{}) if isinstance(v,dict) else {}
    sources=v.get('sources',[]) if isinstance(v,dict) else []
    return {'page_no':sources[0].get('page') if sources else None,'source_text':sources[0].get('text') if sources else None,
      'confidence':v.get('confidence') if isinstance(v,dict) else None}


def field_view(code,item):
    item=copy.deepcopy(item);confidence=item.get('confidence')
    item.update(code=code,label=LABELS.get(code,code))
    item['needs_review']=bool(item.get('needs_review') or item.get('value') is None or not item.get('source_text')
      or not item.get('page_no') or confidence is None or confidence<.85)
    item['excel_supported']=code not in MAPPING or bool(MAPPING[code]['excel_named_range'])
    if code in RULES_MAPPING:item['excel_supported']=False
    if code in PURCHASE_MAPPING:item['excel_supported']=bool(PURCHASE_MAPPING[code]['excel_named_range'])
    return item

LABELS.update({k:v['label'] for k,v in RULES_MAPPING.items() if k not in LABELS})
LABELS.update({k:v['label'] for k,v in PURCHASE_MAPPING.items()})


class Workspace:
    def __init__(self,demo=None):
        load_env();self.demo=flag('DEMO_MODE',False) if demo is None else demo
        self.remote=not self.demo or flag('DEMO_WRITE_DB')
        self.reads=0;self.writes=0;self.ai_calls=0;self.lock=threading.RLock();self.cases=[];self.raw={};self.unmatched=[];self.files={}
        self.template=Path(env('CONTRACT_TEMPLATE') or ROOT.parent/'契約書ひな形_名前定義版.xlsm')
        self.output=Path(env('WEB_OUTPUT_DIR') or ROOT/'output/web')
        if self.remote:self.refresh()
        else:self.load_demo()
        from web_registry import restore
        if not self.remote:restore(self)
        if not self.remote:
            from web_report import restore as restore_reports
            restore_reports(self)
            from web_rules import restore as restore_rules
            restore_rules(self)
            from web_purchase import restore as restore_purchase
            restore_purchase(self)

    def rpc(self,name,p=None,write=False):
        url,key=env('SUPABASE_URL'),env('SUPABASE_SERVICE_ROLE_KEY')
        import re
        if not url or not key or not re.fullmatch(r'https://[a-z0-9]+\.supabase\.co/?',url):
            raise StoreError('Supabaseの接続設定を確認してください。')
        req=urllib.request.Request(url.rstrip('/')+'/rest/v1/'+name,data=canonical({'p':p or {}}).encode('utf8'),
            headers={'apikey':key,'Authorization':'Bearer '+key,'Content-Type':'application/json'})
        if write:self.writes+=1
        else:self.reads+=1
        try:
            with urllib.request.build_opener(NoRedirect()).open(req,timeout=60) as r:return json.load(r)
        except (OSError,ValueError):raise StoreError('DB操作を完了できませんでした。再読込して差分・根拠を確認してください。') from None

    def load_demo(self):
        registry=env('DEMO_REGISTRY_JSON') or ROOT.parent/'verification_names/integrated_claude_private.json'
        report=env('DEMO_REPORT_JSON')
        evidence=read_json(env('DEMO_EVIDENCE_JSON')) if env('DEMO_EVIDENCE_JSON') else {}
        if not report:raise StoreError('DEMO_REPORT_JSONに解析済み重調JSONを設定してください。')
        r=read_json(registry);h=read_json(report)
        if not r.get('専有部分'):raise StoreError('DEMO_REGISTRY_JSONには既存のClaude形式の中間JSONを指定してください。')
        hv=lambda k:h['fields'].get(k,{}).get('value')
        house=evidence.get('unit',{}).get('house_number',{}).get('value') or env('DEMO_HOUSE_NUMBER')
        if not house:raise StoreError('謄本の根拠JSONまたは完全な家屋番号を設定してください。')
        converted=convert(r,house_number=house,building_name=hv('building_name'))
        unit=converted['unit'];building=converted['building']
        for code in ('management_fee','repair_reserve_fee'):
            unit[code]=hv(code)
        fields=[field_view(x['field_code'],{**x,**evidence_item(evidence,x['field_code'])}) for x in converted['values']]
        report_fields=[field_view(k,v) for k,v in h['fields'].items()]
        docs=[{'id':'demo-registry','version_id':'demo-registry-v1','type':'registry','filename':Path(registry).name,'version':1,'date':None,'fields':fields},
              {'id':'demo-report','version_id':'demo-report-v1','type':'important_report','filename':h.get('source',{}).get('original_filename','重調解析済みJSON'),'file_hash':h.get('source',{}).get('file_hash'),'version':1,'date':None,'fields':report_fields}]
        diffs=[]
        for code,new in [('management_fee',(hv('management_fee') or 0)+1500),('current_owner_name','デモ用変更候補（実在しない法人）')]:
            old=unit.get(code)
            diffs.append({'id':'demo-'+code,'code':code,'label':LABELS[code],'old_value':old,'new_value':new,
              'review_status':'unreviewed','demo':True,'can_adopt':True,'source':'操作確認用の差分（原本の記載ではありません）'})
        c={'id':'demo-case','unit_id':'demo-unit','building_name':hv('building_name'),'unit_name':hv('unit_name'),
          'address':hv('display_address') or building.get('registry_location'),'owner':unit.get('current_owner_name'),
          'area':unit.get('registered_area'),'status':'デモ確認中','updated_at':None,'documents':docs,'diffs':diffs}
        self.cases=[c];self.raw={c['id']:{'registry':r,'unit':unit,'building':building,'report_fields':h['fields']}}

    def refresh(self):
        if not self.remote:return
        s=self.rpc('rpc/web_workspace_snapshot')
        self.cases=[];self.raw={};self.unmatched=[]
        units={x['unit_id']:x for x in s['units']};buildings={x['building_id']:x for x in s['buildings']}
        docs={x['document_id']:x for x in s['documents']};versions={x['document_version_id']:x for x in s['versions']}
        rows=s['cases'][:]
        for uid in units:
            if not any(c.get('unit_id')==uid for c in rows):rows.append({'case_id':'unit-'+uid,'unit_id':uid,'case_status':'案件未登録'})
        for case in rows:
            unit=units.get(case.get('unit_id'))
            if not unit:continue
            b=buildings[unit['building_id']];documents=[]
            for vid,v in versions.items():
                d=docs[v['document_id']]
                if d.get('unit_id')!=unit['unit_id'] and not (d['document_type']=='management_rules' and d.get('building_id')==unit['building_id'] and d.get('unit_id') is None):continue
                values=[]
                for e in s['values']:
                    if e['document_version_id']!=vid:continue
                    meta=(v.get('important_raw_json') or v.get('management_rules_raw_json') or {}).get('fields',{}).get(e['field_code'],{})
                    provenance=e.get('registry_provenance') or {}
                    values.append(field_view(e['field_code'],{**meta,**e,'needs_review':bool(meta.get('needs_review') or provenance.get('needs_review'))}))
                    if d['document_type']=='management_rules':values[-1]['excel_supported']=False
                documents.append({'id':d['document_id'],'version_id':vid,'type':d['document_type'],'filename':v.get('original_filename'),
                  'version':v['version_no'],'date':v.get('as_of_date') or v.get('uploaded_at'),'fields':values})
            diffs=[]
            for x in s['diffs']:
                document=docs[x['document_id']]
                if case.get('purchase_values') is not None and x.get('affected_field')!='purchase:'+case['case_id']:continue
                if case.get('purchase_values') is None and (x.get('affected_field') or '').startswith('purchase:'):continue
                if document.get('unit_id')!=unit['unit_id'] and not (document['document_type']=='management_rules' and document.get('building_id')==unit['building_id'] and document.get('unit_id') is None):continue
                e=next((e for e in s['values'] if e['document_version_id']==x['new_version_id'] and e['field_code']==x['field_code']),{})
                version=versions.get(x['new_version_id'],{})
                meta=(version.get('important_raw_json') or version.get('management_rules_raw_json') or {}).get('fields',{}).get(x['field_code'],{})
                provenance=e.get('registry_provenance') or {}
                trusted=not field_view(x['field_code'],{**meta,**e,'needs_review':bool(meta.get('needs_review') or provenance.get('needs_review'))})['needs_review']
                diffs.append({**x,'id':x['diff_id'],'code':x['field_code'],'label':LABELS.get(x['field_code'],x['field_code']),
                  'can_adopt':trusted,'document_type':document['document_type'],'source':versions.get(x['new_version_id'],{}).get('original_filename')})
            c={'id':case['case_id'],'unit_id':unit['unit_id'],'building_id':unit['building_id'],'building_name':b.get('building_name'),'unit_name':unit.get('unit_name'),
              'address':b.get('display_address') or b.get('registry_location'),'owner':unit.get('current_owner_name'),'area':unit.get('registered_area'),
              'status':case.get('case_status') or '確認中','updated_at':case.get('updated_at') or unit.get('updated_at'),'documents':documents,'diffs':diffs}
            self.cases.append(c)
            imports=[i for i in s['imports'] if docs[versions[i['document_version_id']]['document_id']].get('unit_id')==unit['unit_id']]
            imports.sort(key=lambda i:(i.get('created_at') or versions[i['document_version_id']].get('uploaded_at') or '',i['content_hash']))
            # Start from the original registry, then overlay explicitly adopted master values.
            registry=copy.deepcopy(imports[0]['raw_json']) if imports else None
            report_fields={}
            for doc in documents:
                if doc['type']=='important_report':
                    for f in doc['fields']:
                        if f.get('approved'):report_fields[f['code']]=f
            if registry and registry.get('format')=='normalized_registry_v1':
                from web_registry import case_from
                normalized=registry['integrated']
                safety=case_from(normalized,c['id'])
                c.update(property_type=safety['property_type'],generation_blocked=safety['generation_blocked'])
                self.raw[c['id']]={'uploaded':normalized,'registered':True,'unit':unit,'building':b,'report_fields':report_fields}
            else:self.raw[c['id']]={'registry':registry,'unit':unit,'building':b,'report_fields':report_fields}
            from web_purchase import apply_snapshot
            apply_snapshot(self,c,case,s)
        for d in docs.values():
            if not d.get('unit_id') and not (d['document_type']=='management_rules' and d.get('building_id')):self.unmatched.append({'type':d['document_type'],'title':d.get('title') or '住戸未照合の資料'})
        from web_registry import restore
        restore(self)
        from web_purchase import restore as restore_purchase
        restore_purchase(self)

    def public(self):
        result=copy.deepcopy(self.cases)
        for c in result:
            from web_rules import conflicts
            from web_documents import pending_names
            c['source_conflicts']=conflicts(c)
            from web_choices import decorate
            decorate(self,c)
            c['pending_documents']=pending_names(self,c['id'])
            c['review_count']=sum(f['needs_review'] for d in c['documents'] for f in d['fields'])
            c['diff_count']=sum(d['review_status']=='unreviewed' for d in c['diffs'])
        return {'mode':'demo' if self.demo else 'live','db_write_enabled':self.remote,'cases':result,'unmatched':self.unmatched,
          'usage':{'ai_calls':self.ai_calls,'db_reads':self.reads,'db_writes':self.writes}}

    def case(self,cid):
        case=next((c for c in self.cases if c['id']==cid),None)
        if case is None:raise StoreError('案件が見つかりません。')
        return case

    def decide(self,cid,did,action):
        with self.lock:
            c=self.case(cid);diff=next((d for d in c['diffs'] if d['id']==did),None)
            if not diff or action not in ('adopt','hold'):raise StoreError('差分の操作を確認してください。')
            if diff['review_status'] in ('applied','ignored'):raise StoreError('処理済みの差分です。')
            if action=='adopt' and not diff.get('can_adopt'):raise StoreError('根拠が不明なため採用できません。原本確認が必要です。')
            if self.remote:
                if (diff.get('affected_field') or '').startswith('purchase:'):self.rpc('rpc/web_review_purchase_diff',{'diff_id':did,'case_id':cid,'action':action},write=True)
                elif diff.get('document_type')=='management_rules':self.rpc('rpc/review_management_rules_diff',{'diff_id':did,'case_id':cid,'action':action},write=True)
                else:self.rpc('rpc/web_review_diff',{'diff_id':did,'unit_id':c['unit_id'],'action':action},write=True)
                self.refresh()
            else:
                if action=='adopt':
                    unit=self.raw[cid]['unit'];code=diff['code']
                    if unit.get(code)!=diff['old_value']:raise StoreError('確定値が変更されています。差分を確認し直してください。')
                    unit[code]=copy.deepcopy(diff['new_value'])
                    if code=='current_owner_name':c['owner']=diff['new_value']
                diff['review_status']='applied' if action=='adopt' else 'reviewed'
                c['updated_at']=dt.datetime.now(dt.timezone.utc).isoformat()
            return self.public()

    def generate(self,cid):
        with self.lock:
            from web_documents import pending_names
            if self.remote and pending_names(self,cid):raise StoreError('未保存の追加資料があります。物件・資料で保存を完了してください。')
            c=self.case(cid);raw=self.raw[cid]
            if 'purchase' in raw:
                from web_purchase import generate
                return generate(self,cid)
            if 'uploaded' in raw:
                from web_registry import generate
                return generate(self,cid)
            if not raw['registry']:raise StoreError('謄本の中間JSONがないため生成できません。')
            data=copy.deepcopy(raw['registry']);u=raw['unit']
            if not data.get('専有部分'):raise StoreError('区分マンション用の謄本JSONを確認してください。')
            for col,key in [('current_owner_name','氏名'),('current_owner_address','住所')]:
                if u.get(col) is not None:data.setdefault('所有者',{})[key]=u[col]
            if u.get('registered_area') is not None:data['専有部分']['床面積']=u['registered_area']
            for col,key in [('land_right_type','権利の種類'),('land_right_numerator','持分_分子'),('land_right_denominator','持分_分母')]:
                if u.get(col) is not None:
                    for land in data.get('土地',[]):land[key]=u[col]
            token=uuid.uuid4().hex;folder=self.output/token;folder.mkdir(parents=True)
            intermediate=folder/'registry_private.json';intermediate.write_text(canonical(data),encoding='utf8')
            registry_out=folder/'registry.xlsm'
            proc=subprocess.run([sys.executable,'-X','utf8',str(SCRIPTS/'fill_tohon.py'),str(intermediate),str(self.template),'-o',str(registry_out)],
                capture_output=True,text=True,encoding='utf8',timeout=120,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            if proc.returncode not in (0,1) or not registry_out.exists():raise StoreError('Excelを生成できませんでした。ひな形を確認してください。')
            from web_excel import preserve_template
            preserved=folder/'registry_preserved.xlsm'
            registry_written=preserve_template(self.template,registry_out,preserved)
            fields=copy.deepcopy(raw['report_fields'])
            if not self.remote:
                # Only explicitly configured demo baseline fields may be exported.
                approved=set((env('DEMO_APPROVED_FIELDS') or '').split(','))
                for code,f in fields.items():f['approved']=code in approved
            final=folder/('契約書_自動入力_DEMO.xlsm' if self.demo else '契約書_自動入力済.xlsm')
            report=write_named_excel(preserved,final,fields)
            if not report['output']:
                import shutil
                shutil.copyfile(preserved,final)
            from openpyxl import load_workbook
            check=load_workbook(final,read_only=True,data_only=False)
            try:
                if '基本入力' not in check.sheetnames:raise StoreError('Excel再読込に失敗しました。')
            finally:check.close()
            self.files[token]=final
            warnings=[s.strip().lstrip('!').strip() for s in proc.stdout.splitlines() if '!' in s]
            return {'download':'/download/'+token,'filename':final.name,'report_written':len(report['written']),
              'registry_written':len(registry_written),
              'warnings':warnings,'unsupported':[v['label'] for v in MAPPING.values() if not v['excel_named_range']],
              'message':'生成・再読込確認が完了しました。要確認事項と出力内容を確認してください。'}
